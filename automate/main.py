#!/usr/bin/env python3
import pandas as pd
from openpyxl import load_workbook

def main():
    # Input and output file names – change these as needed.
    input_file = 'input.xlsx'
    output_file = 'output_validated.xlsx'
    
    # Allowed values for the last six columns (case sensitive)
    allowed_values = {
        "Functional Area": [
            "CRIT", "CRIT - Data Management", "CRIT - Data Governance", 
            "CRIT - Regulatory Reporting", "CRIT - Portfolio Reporting", "CRIT - Transformation"
        ],
        "Project Category": [
            "Data Infrastructure", "Monitoring & Insights", "Analytics / Strategy Development", 
            "GDA Related", "Trainings and Team Meeting"
        ],
        "Complexity": ["H", "M", "L"],
        "Novelity": ["BAU repetitive", "One time repetitive", "New one time"],
        "Output Type": [
            "Core production work", "Ad-hoc long-term projects", "Ad-hoc short-term projects", 
            "Business Management", "Administration", "Trainings/L&D activities", "Others"
        ],
        "Impact type": ["Customer Experience", "Financial impact", "Insights", "Risk reduction", "Others"]
    }
    
    # Open the workbook to get all sheet names.
    wb = load_workbook(input_file)
    sheet_names = wb.sheetnames

    # The first sheet 'Home' contains the employee names in column F.
    # Assuming row 1 contains headers and row 2 (cell F2) is the header for employee names,
    # then from row 3 downward (index 1 in the DataFrame) are the employee names.
    home_df = pd.read_excel(input_file, sheet_name='Home')
    employee_names = home_df.iloc[1:, 5].dropna().tolist()  # column index 5 corresponds to column F
    employee_names = [str(name) for name in employee_names]

    # Prepare storage for violations and monthly reports.
    violations = []  # Each violation will be recorded as a dict with keys: Employee, Violation Type, Location.
    monthly_reports = {}  # key: employee name, value: monthly summary DataFrame.

    # Dictionary to track a project’s first recorded start date.
    project_start_dates = {}  # key: project name, value: start date

    # Process each employee sheet (if found in the workbook)
    for emp in employee_names:
        if emp in sheet_names:
            df = pd.read_excel(input_file, sheet_name=emp)
            # Expected columns (order matters):
            expected_columns = [
                "Status Date (Every Friday)", "Main project", "Name of the Project", 
                "Project Key Milestones", "TM", "Start Date", "Completion Date % of Completion", 
                "Status", "Weekly Time Spent(Hrs)", "Projected hours (Based on the Project: End to End implementation)", 
                "Functional Area", "Project Category", "Complexity", "Novelity", "Output Type", "Impact type"
            ]
            # (You can add a check here to ensure df.columns match expected_columns.)
            
            # Loop over each row to run validations.
            for idx, row in df.iterrows():
                # Compute an Excel-style row number (assuming header is row 1)
                excel_row = idx + 2  
                
                # --- Validation 1: Check that each of the last six columns has an allowed value.
                for col, allowed in allowed_values.items():
                    value = row[col]
                    if value not in allowed:
                        violations.append({
                            "Employee": emp,
                            "Violation Type": f"Invalid value in {col}: '{value}'",
                            "Location": f"Sheet: {emp}, Row: {excel_row}"
                        })
                
                # --- Validation 3: Start Date consistency.
                project_name = row["Name of the Project"]
                start_date = row["Start Date"]
                if project_name not in project_start_dates:
                    project_start_dates[project_name] = start_date
                else:
                    if start_date != project_start_dates[project_name]:
                        violations.append({
                            "Employee": emp,
                            "Violation Type": f"Start date changed for project '{project_name}'",
                            "Location": f"Sheet: {emp}, Row: {excel_row}"
                        })
            
            # Convert the "Status Date (Every Friday)" column to datetime.
            df["Status Date (Every Friday)"] = pd.to_datetime(df["Status Date (Every Friday)"], errors='coerce')
            # Create a new column for month (as a period) from the status date.
            df["Month"] = df["Status Date (Every Friday)"].dt.to_period("M")
            
            # Calculate work hours and PTO hours.
            # (If "Main project" equals "PTO" then treat the hours as PTO hours.)
            df["Work Hours"] = df.apply(lambda r: r["Weekly Time Spent(Hrs)"] if r["Main project"] != "PTO" else 0, axis=1)
            df["PTO Hours"] = df.apply(lambda r: r["Weekly Time Spent(Hrs)"] if r["Main project"] == "PTO" else 0, axis=1)
            
            # Create monthly summary for this employee.
            monthly_summary = df.groupby("Month").agg({"Work Hours": "sum", "PTO Hours": "sum"}).reset_index()
            monthly_reports[emp] = monthly_summary

            # --- Validation 2: Check that the employee has worked at least 40 hours in each week.
            # Group rows by week. We assume that the "Status Date" is always a Friday,
            # so we use a year-week grouping based on the date.
            df["Week"] = df["Status Date (Every Friday)"].dt.strftime("%Y-%U")
            weekly_summary = df.groupby("Week").agg({"Work Hours": "sum"}).reset_index()
            for _, week_row in weekly_summary.iterrows():
                if week_row["Work Hours"] < 40:
                    violations.append({
                        "Employee": emp,
                        "Violation Type": f"Weekly work hours less than 40 (Total: {week_row['Work Hours']}) in week {week_row['Week']}",
                        "Location": f"Sheet: {emp}, Week: {week_row['Week']}"
                    })
        else:
            print(f"Warning: Sheet for employee '{emp}' not found in the workbook.")
    
    # Write out a new Excel file with:
    # - The original sheets (Home and each employee)
    # - New monthly report sheets for each employee (named <Employee>_Report)
    # - A new "Violations" sheet listing all violations.
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        # Write all original sheets.
        for sheet in sheet_names:
            sheet_df = pd.read_excel(input_file, sheet_name=sheet)
            sheet_df.to_excel(writer, sheet_name=sheet, index=False)
        
        # Write monthly report sheets for each employee.
        for emp, report_df in monthly_reports.items():
            # Limit sheet name length (Excel sheet names can have at most 31 characters).
            sheet_name = f"{emp}_Report"
            if len(sheet_name) > 31:
                sheet_name = sheet_name[:31]
            report_df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        # Write the violations sheet.
        violations_df = pd.DataFrame(violations)
        violations_df.to_excel(writer, sheet_name="Violations", index=False)
    
    print("Validation and report generation completed. Output saved to", output_file)

if __name__ == "__main__":
    main()